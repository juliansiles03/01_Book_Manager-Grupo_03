import csv

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select

from book_manager.database.connection import ConexionDB
from book_manager.models.models import (
    MonedaModelo,
    PrecioModelo,
)


COLUMNAS_REPORTE: tuple[str, ...] = (
    "Libro",
    "Precio interno",
    "Precio web",
    "Diferencia",
    "Fecha de extracción",
)


def generar_reporte_precios(
    ruta_extraccion: str | Path,
    ruta_reporte: str | Path,
) -> int:
    """Genera un Excel con los precios internos y los extraídos."""
    origen = Path(ruta_extraccion)
    destino = Path(ruta_reporte)

    if not origen.is_file():
        raise FileNotFoundError(
            f"No existe la extracción: {origen}"
        )

    conexion = ConexionDB()

    try:
        with conexion.transaccion() as sesion:
            registros = sesion.execute(
                select(
                    PrecioModelo.libro_id,
                    PrecioModelo.valor,
                )
                .join(
                    MonedaModelo,
                    PrecioModelo.moneda_id
                    == MonedaModelo.id,
                )
                .where(
                    MonedaModelo.codigo == "ARS"
                )
            ).all()

        precios_internos: dict[int, Decimal] = {
            int(libro_id): Decimal(str(valor))
            for libro_id, valor in registros
        }
    finally:
        conexion.cerrar()

    libro_excel = Workbook()
    hoja = libro_excel.active
    hoja.title = "Comparación"
    hoja.append(COLUMNAS_REPORTE)

    cantidad = 0
    zona_argentina = ZoneInfo(
        "America/Argentina/Buenos_Aires"
    )

    with origen.open(
        "r",
        encoding="utf-8",
        newline="",
    ) as archivo:
        for fila in csv.DictReader(archivo):
            libro_id = int(fila["libro_id"])

            if libro_id not in precios_internos:
                raise ValueError(
                    "Falta un precio interno ARS para "
                    f"el libro ID {libro_id}."
                )

            precio_interno = precios_internos[
                libro_id
            ]
            precio_web = Decimal(
                fila["precio_web"]
            )
            diferencia = (
                precio_web - precio_interno
            )

            fecha = datetime.fromisoformat(
                fila["fecha_extraccion"]
            )

            if fecha.tzinfo is not None:
                fecha = fecha.astimezone(
                    zona_argentina
                ).replace(tzinfo=None)

            numero_fila = hoja.max_row + 1
            hoja.append(
                (
                    fila["titulo"],
                    float(precio_interno),
                    float(precio_web),
                    float(diferencia),
                    fecha,
                )
            )

            for columna in ("B", "C", "D"):
                hoja[
                    f"{columna}{numero_fila}"
                ].number_format = '#,##0.00'

            hoja[
                f"E{numero_fila}"
            ].number_format = "dd/mm/yyyy hh:mm"

            cantidad += 1

    for celda in hoja[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="244062",
        )
        celda.alignment = Alignment(
            horizontal="center",
        )

    for columna, ancho in {
        "A": 43,
        "B": 18,
        "C": 18,
        "D": 18,
        "E": 23,
    }.items():
        hoja.column_dimensions[
            columna
        ].width = ancho

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    hoja.sheet_view.showGridLines = False

    destino.parent.mkdir(
        parents=True,
        exist_ok=True,
    )
    libro_excel.save(destino)

    return cantidad
